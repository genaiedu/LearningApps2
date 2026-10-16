"""Extract the original MG pH matrix; pass the unchanged Goto workbook as argv[1]."""
import sys,json,hashlib
from pathlib import Path
import openpyxl
source=Path(sys.argv[1]);sheet=openpyxl.load_workbook(source,data_only=True).worksheets[0]
curves=[]
for col in range(2,19):
    pH=sheet.cell(1,col).value
    points=[]
    for row in range(2,603):
        w,e=sheet.cell(row,1).value,sheet.cell(row,col).value
        if isinstance(w,(int,float)) and isinstance(e,(int,float)) and 200<=w<=800:points.append([w,e])
    assert len(points)>=500
    curves.append(dict(pH=pH,points=points))
record=dict(id='PH-MG',name='Malachitgrün · pH-Messreihe',curves=curves,unit='absorbance',source='https://data.mendeley.com/datasets/jp2j9hz6v9/1',dataURL='https://data.mendeley.com/public-files/datasets/jp2j9hz6v9/files/db784688-fd53-42d1-b166-7dd18d3d5fbd/file_downloaded',sheet='Sheet1',credit='Satoru Goto (2026)',rights='CC BY 4.0',sha256=hashlib.sha256(source.read_bytes()).hexdigest(),note='Unveränderte Original-Extinktion; Konzentration nicht ausgewiesen. Die starke Säureprotonierung und die basische Carbinolbildung sind verschiedene Prozesse.')
(Path(__file__).resolve().parents[1]/'data/uvvis-indicator-extra.json').write_text(json.dumps(record,separators=(',',':'),ensure_ascii=False)+'\n')
print('Extracted',len(curves),'pH spectra,',len(curves[0]['points']),'wavelengths each')
