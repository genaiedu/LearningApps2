"""Synchronize the static assessment chapters of the approved working copy.

Preserve the existing DOCX, especially all teaching units, figures and GKL
rubrics. Use its current HTML as the single source for the added prose and
tables. Interactive controls are intentionally excluded. Render with the
bundled documents renderer to create the matching PDF before publication.
"""
from pathlib import Path
from copy import deepcopy
from urllib.parse import urljoin
import argparse
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import html

REPO = Path(__file__).resolve().parents[1]
STEM = 'Curriculum_Chemie_Arbeitskopie_Thomaeum_Sek_I_und_II'
BASE_URL = 'https://genaiedu.github.io/LearningApps2/'
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source', type=Path, default=REPO/'downloads'/f'{STEM}.docx')
parser.add_argument('--output', type=Path, required=True)
args = parser.parse_args()
doc = Document(args.source)
source = html.fromstring((REPO/'chemie-curriculum-arbeitskopie.html').read_text())
body = doc._element.body
bookmark_id = max(int(b.get(qn('w:id'))) for b in doc._element.xpath('//w:bookmarkStart')) + 1


def xmltext(el):
    return ''.join(el.xpath('.//w:t/text()'))


def style(el):
    return el.xpath('./w:pPr/w:pStyle/@w:val')


def optional_hyphens(element):
    """Native discretionary hyphens avoid arbitrary breaks in compound words."""
    words={'Darstellungsebenen':'Darstellungs|ebenen',
           'Stoffmengenrechnung':'Stoffmengen|rechnung',
           'Quellenauswertung':'Quellen|auswertung',
           'Notenpunkte':'Noten|punkte'}
    for node in list(element.xpath('.//w:t')):
        value=node.text or ''
        for word,parts in words.items():
            if word not in value:
                continue
            start=value.index(word)+len(parts.split('|')[0])
            node.text=value[:start]
            hyphen=OxmlElement('w:softHyphen')
            rest=OxmlElement('w:t')
            rest.text=value[start:]
            node.addnext(hyphen)
            hyphen.addnext(rest)
            optional_hyphens(rest.getparent())
            break


def patch_text(el, replacements):
    """Replace spans across run boundaries without discarding links or styles."""
    for old, new in replacements:
        nodes = el.xpath('.//w:t')
        full = ''.join(n.text or '' for n in nodes)
        for match in reversed(list(re.finditer(re.escape(old), full))):
            position = 0
            first = True
            for node in nodes:
                value = node.text or ''
                lo, hi = position, position + len(value)
                position = hi
                if hi <= match.start() or lo >= match.end():
                    continue
                left, right = max(0, match.start()-lo), min(len(value), match.end()-lo)
                node.text = value[:left] + (new if first else '') + value[right:]
                first = False


def bookmark(p, name):
    global bookmark_id
    name = name.replace('-', '_')
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))
    p._p.insert(1, start)
    p._p.append(end)
    bookmark_id += 1


def fresh_list_id():
    numbering=doc.part.numbering_part.element
    num_id=max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num')))+1
    style_num=doc.styles['List Number'].element.xpath('.//w:numId/@w:val')[0]
    abstract=numbering.xpath(f'./w:num[@w:numId="{style_num}"]/w:abstractNumId/@w:val')[0]
    num=OxmlElement('w:num')
    num.set(qn('w:numId'),str(num_id))
    base=OxmlElement('w:abstractNumId')
    base.set(qn('w:val'),abstract)
    num.append(base)
    override=OxmlElement('w:lvlOverride')
    override.set(qn('w:ilvl'),'0')
    start=OxmlElement('w:startOverride')
    start.set(qn('w:val'),'1')
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def number_paragraph(element,identity):
    props=element.get_or_add_pPr()
    for old in props.findall(qn('w:numPr')):
        props.remove(old)
    numbering=OxmlElement('w:numPr')
    level=OxmlElement('w:ilvl')
    level.set(qn('w:val'),'0')
    num_id=OxmlElement('w:numId')
    num_id.set(qn('w:val'),str(identity))
    numbering.extend([level,num_id])
    props.append(numbering)


def inline(p, element, bold=False, italic=False):
    if element.text:
        run = p.add_run(element.text)
        run.bold, run.italic = bold, italic
    for child in element:
        if not isinstance(child.tag, str):
            continue
        if child.tag == 'br':
            p.add_run().add_break()
        elif child.tag == 'a':
            link = OxmlElement('w:hyperlink')
            target = child.get('href', '')
            if 'downloads/' in target or target=='#links':
                continue
            if target.startswith('#'):
                link.set(qn('w:anchor'), target[1:].replace('-', '_'))
            else:
                link.set(qn('r:id'), doc.part.relate_to(urljoin(BASE_URL, target), RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            rp = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '174D72')
            rp.append(color)
            run.append(rp)
            value = OxmlElement('w:t')
            value.text = child.text_content()
            run.append(value)
            link.append(run)
            p._p.append(link)
        elif child.tag != 'input':
            inline(p, child, bold or child.tag in ('strong', 'b'), italic or child.tag in ('em', 'i'))
        if child.tail:
            run = p.add_run(child.tail)
            run.bold, run.italic = bold, italic


def addhtml(element, target):
    if not isinstance(element.tag, str):
        return
    classes = element.get('class', '').split()
    if (element.tag in ('script','style','input','button','select','textarea')
        or element.get('id') == 'klausurassistent'
        or any(c in classes for c in ('assessment-builder','klausur-tool'))):
        return
    tag = element.tag
    if 'actions' in classes:
        return
    if element.get('id') == 'muendliche-pruefung-vorlagen':
        p = target.add_paragraph('Vorlage für die Prüfungsunterlagen',style='Heading 4')
        p.runs[0].bold=True
        p = target.add_paragraph('Die vollständige Vorlage für beide Prüfungsteile steht in Anhang A. Sie umfasst Prüfungsdaten, unterrichtliche Voraussetzungen, Aufgabenblatt, Erwartungshorizont, Gesprächsplanung, Bewertungsbegründung und Checkliste. Die Anforderungsbereiche werden ausschließlich nach Abschnitt 13.2 zugeordnet. Die Vorlage dient der Vorbereitung und Dokumentation; die ausgefüllten Bewertungsunterlagen werden nicht an den Prüfling ausgegeben.')
        return
    if tag=='a' and (element.get('href')=='#links' or 'downloads/' in element.get('href','')):
        return
    if tag in ('h2','h3','h4','h5','p','summary'):
        styles = {'h2':'Heading 1','h3':'Heading 2','h4':'Heading 3','h5':'Heading 4','summary':'Heading 4'}
        p = target.add_paragraph(style=styles.get(tag, 'Normal'))
        inline(p, element)
        p.paragraph_format.space_after = Pt(7)
        if tag != 'p':
            p.paragraph_format.keep_with_next = True
        if tag in ('h4','h5','summary'):
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0,0,0)
        if tag == 'h2':
            p.paragraph_format.page_break_before = True
        if element.get('id') == 'klausuren-2':
            p.paragraph_format.page_break_before = True
        if element.get('id'):
            bookmark(p, element.get('id'))
        if 'Häkchen dienen nur der Vorbereitung' in p.text:
            p.text = 'Leitgedanke: So konkret wie nötig – so offen wie möglich.'
            p.runs[0].bold = True
        if 'Schulinterne Verweise und die Dokumentfassungen stehen' in p.text:
            p.text = 'Die folgenden Quellenverweise führen zu offiziellen NRW-Veröffentlichungen. Schulinterne Verweise stehen unter „Links und Onlinefassungen“.'
    elif tag in ('ul','ol'):
        num_id=fresh_list_id() if tag=='ol' else None
        for li in element:
            if li.tag != 'li':
                continue
            check = 'assessment-checklist' in classes
            p = target.add_paragraph(style='Normal' if check else ('List Number' if tag=='ol' else 'List Bullet'))
            if num_id is not None:
                number_paragraph(p._p,num_id)
            if check:
                p.add_run('☐ ').font.name = 'DejaVu Sans'
            inline(p, li)
            p.paragraph_format.space_after = Pt(5)
    elif tag == 'table':
        rows = element.xpath('./thead/tr|./tbody/tr|./tr')
        if not rows:
            return
        caption = element.find('caption')
        if caption is not None:
            p = target.add_paragraph(caption.text_content(), style='Caption')
            p.paragraph_format.keep_with_next = True
        table = target.add_table(rows=0, cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        width = doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
        count = len(rows[0])
        weights = {2:[.25,.75],3:[.22,.40,.38],4:[.22,.26,.26,.26],5:[.15,.24,.21,.20,.20]}[count]
        if count==3 and 'Notenpunkte' in rows[0].text_content():
            weights = [.20,.16,.64]
        for column, weight in zip(table.columns,weights):
            column.width = int(width*weight)
        for i, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, original, weight in zip(cells,row,weights):
                cell.width = int(width*weight)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for side in ('top','left','bottom','right'):
                    edge = OxmlElement('w:'+side)
                    for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:
                        edge.set(qn('w:'+key),value)
                    borders.append(edge)
                cp.append(borders)
                margins = OxmlElement('w:tcMar')
                for side in ('top','left','bottom','right'):
                    margin = OxmlElement('w:'+side)
                    margin.set(qn('w:w'),'100')
                    margin.set(qn('w:type'),'dxa')
                    margins.append(margin)
                cp.append(margins)
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'),'E9EEF0' if i==0 else 'FFFFFF')
                cp.append(shade)
                inline(cell.paragraphs[0], original)
                optional_hyphens(cell._tc)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = False
                    p.paragraph_format.line_spacing = 1.05
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(0,0,0)
                        if i==0 or original.tag=='th':
                            run.bold = True
            rowprops = table.rows[i]._tr.get_or_add_trPr()
            rowprops.append(OxmlElement('w:cantSplit'))
            if i==0:
                rowprops.append(OxmlElement('w:tblHeader'))
        target.add_paragraph().paragraph_format.space_after = Pt(3)
    elif tag == 'a':
        p = target.add_paragraph()
        wrapper = html.Element('span')
        wrapper.append(deepcopy(element))
        inline(p,wrapper)
    elif tag == 'div' and not any(c.tag in ('div','p','h2','h3','h4','h5','ul','ol','table','details') for c in element):
        p = target.add_paragraph()
        inline(p,element)
        if element.get('id'):
            bookmark(p,element.get('id'))
    else:
        if element.get('id'):
            p = target.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = Pt(1)
            bookmark(p,element.get('id'))
        for child in element:
            addhtml(child,target)
        # Text-only divs are notes, not decorative boxes in print.
        if not len(element) or (element.text and element.text.strip()):
            if element.text and element.text.strip():
                p = target.add_paragraph()
                inline(p,element)


def generated(section_id):
    temporary = Document()
    addhtml(source.get_element_by_id(section_id).getparent(),temporary)
    return [deepcopy(n) for n in temporary._element.body if n.tag != qn('w:sectPr')]


elements = list(body)
# Strip previously generated appendices when refreshing an existing export.
filtered=[]
in_appendix=False
for e in elements:
    if style(e)==['Heading1']:
        in_appendix=xmltext(e).startswith('Anhang ')
    if not in_appendix:
        filtered.append(e)
elements=filtered
starts = [i for i,e in enumerate(elements) if style(e)==['Heading1'] and re.match(r'^\d+ ',xmltext(e))]
assert len(starts)==18
chapters = {}
prefix = [e for e in elements[:starts[0]] if not e.xpath('./w:hyperlink[starts-with(@w:anchor,"anhang_")]')]
section_properties = elements[-1]
assert section_properties.tag==qn('w:sectPr')
for index,start in enumerate(starts):
    end = starts[index+1] if index+1<len(starts) else len(elements)-1
    title = re.sub(r'^\d+ ','',xmltext(elements[start]))
    chapters[title] = elements[start:end]

# Reorder sources to the end, matching the website, and fix retained references.
mapping = {13:18,14:13,15:14,16:15,17:16,18:17}
renumber = '13 Quellen und Geltungsstand' in [xmltext(elements[i]) for i in starts]
for e in elements:
    for p in ([e] if e.tag==qn('w:p') else e.xpath('.//w:p')):
        value = xmltext(p)
        changes = []
        if renumber and style(p) in (['Heading1'],['Heading2']):
            m = re.match(r'^(1[3-8])(?=[. ])',value)
            if m:
                changes.append((m.group(),str(mapping[int(m.group())])))
        elif renumber:
            reference_pattern = r'(?:Kapiteln?|Abschnitt(?:e|en)?)\s+\d+(?:\.\d+)?(?:\s*(?:und|bis|,)\s*\d+(?:\.\d+)?)*'
            for m in re.finditer(reference_pattern,value):
                new = re.sub(r'\b1[3-8]\b',lambda n:str(mapping[int(n.group())]),m.group())
                changes.append((m.group(),new))
            for m in re.finditer(r'(?:in|aus|folgt|gilt|Aufgaben|Präsentationsbogen|Experimentierbogen|Modul|Präsentation|Kursanleitung|Polymerstationen|bis|und|sowie)\s*\(?((?:1[3-8])\.\d+)(?!\d|\.\d)',value):
                number=m.group(1)
                main, dot, sub = number.partition('.')
                if not any(m.start() >= c.start() and m.end() <= c.end() for c in re.finditer(reference_pattern,value)):
                    changes.append((m.group(),m.group().replace(number,str(mapping[int(main)])+dot+sub)))
        patch_text(p,changes)
        patch_text(p,[('Stand 29. September 2026','Stand 4. Oktober 2026')])
        if value=='12.3 Einführungsphase' and style(p)==['Heading2']:
            # The generic chapter precedes the register; this was a pre-existing typo.
            previous = next((i for i,b in enumerate(elements) if b is e),-1)
            if previous < starts[11]:
                patch_text(p,[('12.3 Einführungsphase','11.3 Einführungsphase')])

# Rebuild only the assessment chapters; the existing GKL layout stays intact.
chapters['Klausuren gestalten und bewerten'] = generated('klausurengestaltung')
chapters['Abiturprüfung'] = generated('abiturpruefung')
chapters['Quellen und Geltungsstand'] = generated('quellen')
for e in chapters['Quellen und Geltungsstand']:
    if 'Schulischer Beispielprompt' in xmltext(e):
        patch_text(e,[('Kapitel 18','Kapitel 17')])
gkl_title = next(t for t in chapters if t.startswith('Gleichwertige komplexe Leistungsnachweise'))
gkl = chapters[gkl_title]
gkl[:] = [e for e in gkl if not e.xpath('.//w:bookmarkStart[@w:name="gkl_afb_verweis"]')]
# Insert just the requested short reference, not another AFB definition.
reference = source.get_element_by_id('gkl-afb-verweis')
temporary = Document()
addhtml(reference,temporary)
insert = next(i for i,e in enumerate(gkl) if style(e)==['Heading2'])
gkl[insert:insert] = [deepcopy(temporary.paragraphs[0]._p)]
# The four independent GKL steps must start at 1, not continue an earlier list.
steps=[e for e in gkl if style(e)==['ListNumber']]
assert len(steps)==4
step_id=fresh_list_id()
for e in steps:
    number_paragraph(e,step_id)

ordered = [(h.get('id'),h.text_content()) for h in source.xpath('//main//h2')]
assert len(ordered)==18


def appendix_heading(target, title, name):
    p=target.add_paragraph(title,style='Heading 1')
    p.paragraph_format.page_break_before=True
    bookmark(p,name)


def blank(target,label,lines=1):
    p=target.add_paragraph(label)
    p.runs[0].bold=True
    p.paragraph_format.keep_with_next=True
    for _ in range(lines):
        p=target.add_paragraph('_'*76)
        p.paragraph_format.space_after=Pt(4)
        p.runs[0].font.size=Pt(9)


def copy_table(target,original):
    copied=deepcopy(original)
    for props in copied.xpath('.//w:tcPr'):
        margins=props.find(qn('w:tcMar'))
        if margins is None:
            margins=OxmlElement('w:tcMar')
            props.append(margins)
        for side in ('top','bottom'):
            margin=margins.find(qn('w:'+side))
            if margin is None:
                margin=OxmlElement('w:'+side)
                margins.append(margin)
            margin.set(qn('w:w'),'40')
            margin.set(qn('w:type'),'dxa')
    for props in copied.xpath('.//w:pPr'):
        spacing=props.find(qn('w:spacing'))
        if spacing is None:
            spacing=OxmlElement('w:spacing')
            props.append(spacing)
        spacing.set(qn('w:after'),'20')
    target._element.body.insert(-1,copied)


def endnotes(target,summary):
    p=target.add_paragraph('Bewertung und individuelle Rückmeldung',style='Heading 2')
    p.paragraph_format.page_break_before=True
    target.add_paragraph('Alle Angaben beziehen sich auf die nachgewiesene Leistung der einzelnen Person. Kriterien, Rollen und individuelle Nachweise werden vor Beginn vereinbart. Die Raster konkretisieren die Beispiele in Kapitel 14; sie legen keine neue Gewichtung oder zusätzliche Notendefinition fest.')
    for node in summary:
        target._element.body.insert(-1,deepcopy(node))
    for label,lines in [('Name Kurs Datum und Aufgabe',1),('Vereinbarte Rolle und eigener Beitrag',2),('Teil A Persönlicher Ergebnisbeitrag und Dokumentation mit Beleg',2),('Teil B Fachliches Verständnis und Reflexion mit Beleg',2),('Teil C Rollenleistung und gegebenenfalls eigener Beitrag zur Zusammenarbeit',2),('Bewertungseinheiten A / 20 B / 40 C / 40 Gesamt / 100',1),('Begründete Punktstufe und berücksichtigte Hilfen',2),('Stärken nächste Lernschritte und Vereinbarung',2)]:
        blank(target,label,lines)


appendices={}
temporary=Document()
appendix_heading(temporary,'Anhang A Vorlage zur mündlichen Abiturprüfung','anhang_a')
oral_template=Document(REPO/'downloads/Vorlage_Muendliche_Abiturpruefung_Chemie.docx')
# Import native Word text/tables with isolated styles; no rasterized form pages.
stylemap={s.style_id:'AnhangA_'+s.style_id for s in oral_template.styles}
existing=set(doc.styles._element.xpath('./w:style/@w:styleId'))
for s in oral_template.styles:
    if stylemap[s.style_id] in existing:
        continue
    node=deepcopy(s.element)
    node.set(qn('w:styleId'),stylemap[s.style_id])
    node.attrib.pop(qn('w:default'),None)
    name=node.find(qn('w:name'))
    if name is not None: name.set(qn('w:val'),'Anhang A '+s.name)
    for child in node:
        if child.tag in (qn('w:basedOn'),qn('w:next'),qn('w:link')) and child.get(qn('w:val')) in stylemap:
            child.set(qn('w:val'),stylemap[child.get(qn('w:val'))])
    doc.styles._element.append(node)
for i,original in enumerate(oral_template._element.body):
    if original.tag==qn('w:sectPr') or i==0:
        continue
    node=deepcopy(original)
    for item in node.xpath('.//w:pStyle|.//w:rStyle|.//w:tblStyle'):
        if item.get(qn('w:val')) in stylemap:
            item.set(qn('w:val'),stylemap[item.get(qn('w:val'))])
    for value in node.xpath('.//w:t'):
        if value.text and re.fullmatch(r'_{90,}',value.text): value.text='_'*84
        if value.text: value.text=value.text.replace('Das Ausgabeblatt auf Seite 2','Das folgende Ausgabeblatt')
    for width in node.xpath('.//w:gridCol|.//w:tcW'):
        n=width.get(qn('w:w'))
        if n and n.isdigit(): width.set(qn('w:w'),str(round(int(n)*.89)))
    for width in node.xpath('.//w:tblW[@w:type="dxa"]'):
        n=width.get(qn('w:w'))
        if n and n.isdigit(): width.set(qn('w:w'),str(round(int(n)*.89)))
    temporary._element.body.insert(-1,node)
appendices['anhang_a']=[deepcopy(e) for e in temporary._element.body if e.tag!=qn('w:sectPr')]

temporary=Document()
appendix_heading(temporary,'Anhang B Klausurergebnis und Rückmeldung','anhang_b')
temporary.add_paragraph('Das Notenraster und die Berechnung stehen in Abschnitt 13.6. Dieses Blatt dient der individuellen Dokumentation; es enthält keine zusätzliche Definition der Anforderungsbereiche oder der Notenstufen.')
for label,lines in [('Name Kurs Datum Klausurthema und Lehrkraft',2),('Teilaufgaben Höchstwerte erreichte Bewertungseinheiten und Begründung',8),('Erreichte BE maximal erreichbare BE und Prozentanteil',2),('Zuordnung nach Abschnitt 13.6 Notenpunkte beziehungsweise Note in der EF',2),('Stärken der fachlichen Leistung',3),('Nächste Lernschritte und Vereinbarung',3)]:
    blank(temporary,label,lines)
appendices['anhang_b']=[deepcopy(e) for e in temporary._element.body if e.tag!=qn('w:sectPr')]

# The rubric criteria are copied from the retained GKL chapter, not redefined.
def section_nodes(start,stop):
    i=next(i for i,e in enumerate(gkl) if re.match(start,xmltext(e)))
    j=next(i for i in range(i+1,len(gkl)) if re.match(stop,xmltext(gkl[i])))
    return gkl[i+1:j]

for name,title,start,stop,role_prefix in [
    ('anhang_c','Anhang C GKL Präsentation Bewertungsbogen',r'^14\.6 ',r'^14\.7 ','P'),
    ('anhang_d','Anhang D GKL Experimentieren Bewertungsbogen',r'^14\.7 ',r'^14\.8 ','E')]:
    temporary=Document()
    appendix_heading(temporary,title,name)
    temporary.add_paragraph('Druckvorlage nach Kapitel 14. Individuelle Bewertung A / 20 + B / 40 + C / 40 = 100 BE. Vor Beginn werden Aufgabe, Kriterien, genau ein passendes Rollenmodul und persönliche Nachweise vereinbart.')
    temporary.add_paragraph('Bewertende Lehrkraft: __________________________________________________')
    summary=[]
    in_summary=False
    for e in section_nodes(start,stop):
        copy=deepcopy(e)
        if e.tag==qn('w:tbl') and xmltext(e).startswith('A / 20'):
            in_summary=True
        if in_summary:
            summary.append(copy)
            continue
        # Body rubrics are already blank, native Word tables with room for evidence.
        for b in copy.xpath('.//w:bookmarkStart|.//w:bookmarkEnd'):
            b.getparent().remove(b)
        temporary._element.body.insert(-1,copy)
    p=temporary.add_paragraph('Teil C Rollenmodule',style='Heading 2')
    p.paragraph_format.page_break_before=True
    temporary.add_paragraph('Nur das zuvor vereinbarte Modul ausfüllen. Höchstens 40 BE je Person; individuelle Belege gehören zu jedem Kriterium.')
    for k in range(1,4):
        index=next(i for i,e in enumerate(gkl) if xmltext(e).startswith(role_prefix+str(k)+' '))
        temporary._element.body.insert(-1,deepcopy(gkl[index]))
        table=next(e for e in gkl[index+1:] if e.tag==qn('w:tbl'))
        copy_table(temporary,table)
    endnotes(temporary,summary)
    appendices[name]=[deepcopy(e) for e in temporary._element.body if e.tag!=qn('w:sectPr')]

for e in list(body):
    body.remove(e)
for e in prefix:
    body.append(e)
for section_id,title in ordered:
    if section_id=='quellen':
        for nodes in appendices.values():
            for e in nodes: body.append(e)
    suffix = re.sub(r'^\d+ ','',title)
    assert suffix in chapters,suffix
    for e in chapters[suffix]:
        body.append(e)
body.append(section_properties)

# Update the existing linked table of contents, keeping its page-reference fields.
toc = [e for e in prefix if e.xpath('./w:hyperlink[@w:anchor]')]
toc = [e for e in toc if re.match(r'^\d+ ',xmltext(e))]
assert len(toc)==18
for e,(section_id,title) in zip(toc,ordered):
    link = e.find(qn('w:hyperlink'))
    link.set(qn('w:anchor'),section_id.replace('-','_'))
    values = link.xpath('.//w:t')
    values[0].text=title
    for value in values[1:]:
        value.text=''
    for field in e.findall(qn('w:fldSimple')):
        field.set(qn('w:instr'),f'PAGEREF {section_id.replace("-","_")} \\h')
        field.set(qn('w:dirty'),'true')
        for value in field.iter(qn('w:t')):
            value.text=''
    for keep in e.xpath('./w:pPr/w:keepNext'):
        keep.getparent().remove(keep)

source_toc=toc[-1]
for name,nodes in appendices.items():
    e=deepcopy(source_toc)
    link=e.find(qn('w:hyperlink'))
    link.set(qn('w:anchor'),name)
    value=link.xpath('.//w:t')[0]
    value.text=xmltext(nodes[0])
    for field in e.findall(qn('w:fldSimple')):
        field.set(qn('w:instr'),f'PAGEREF {name} \\h')
    source_toc.addprevious(e)

doc.core_properties.modified = __import__('datetime').datetime(2026,10,4,12,0,0)
doc.core_properties.subject = 'Arbeitskopie mit chemiespezifischen AFB-Konkretisierungen und Unterlagen zur mündlichen Abiturprüfung'
setting = doc.settings.element.find(qn('w:updateFields'))
if setting is None:
    setting=OxmlElement('w:updateFields')
    doc.settings.element.append(setting)
setting.set(qn('w:val'),'true')

# Verify the canonical definitions and both complete grade tables before saving.
full = xmltext(body)
definitions = [
    'Gelerntes im vertrauten Zusammenhang wiedergeben und geübte Verfahren sicher einsetzen.',
    'Bekanntes selbstständig ordnen, erklären und auf vergleichbare neue Situationen übertragen.',
    'Komplexe neue Probleme eigenständig bearbeiten, Methoden auswählen, Lösungen begründen und das Vorgehen reflektieren.'
]
for definition in definitions:
    assert full.count(definition)==1,(definition,full.count(definition))
for forbidden in ('Vorlage erstellen','Vorlage kopieren','Für die Vorlage angenommene Punktstufe','Hier im Fenster öffnen','Schulische Arbeitshilfe, keine neue Prüfungsordnung'):
    assert forbidden not in full,forbidden
tables = source.get_element_by_id('abiturpruefung').getparent().xpath('.//table')
assert len(tables)==2
for table in tables:
    rows=table.xpath('./tbody/tr')
    assert len(rows)==16
    for row in rows:
        assert row[-1].text_content() in full
assert full.rfind('Quellen und Geltungsstand') > full.rfind('Mit KI lernen und Ergebnisse verantwortlich prüfen')
assert full.rfind('Quellen und Geltungsstand') > full.rfind('Anhang D')
for link in doc._element.xpath('//w:hyperlink[@r:id]'):
    assert 'downloads/' not in doc.part.rels[link.get(qn('r:id'))].target_ref
assert not re.search(r'herunterladen|download',full,re.I)
# Drop the unused download relationships as well, not just their visible text.
for rel_id,relationship in list(doc.part.rels.items()):
    if relationship.is_external and 'downloads/' in relationship.target_ref:
        del doc.part.rels[rel_id]
args.output.parent.mkdir(parents=True,exist_ok=True)
doc.save(args.output)
print(f'Saved static working copy: {args.output}')
