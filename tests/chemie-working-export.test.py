"""Static working-copy export checks. Run with Python and python-docx/pypdf/lxml."""
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from lxml import html
from pypdf import PdfReader

REPO = Path(__file__).resolve().parents[1]
STEM = 'Curriculum_Chemie_Arbeitskopie_Thomaeum_Sek_I_und_II'
DEFINITIONS = [
    'Gelerntes im vertrauten Zusammenhang wiedergeben und geübte Verfahren sicher einsetzen.',
    'Bekanntes selbstständig ordnen, erklären und auf vergleichbare neue Situationen übertragen.',
    'Komplexe neue Probleme eigenständig bearbeiten, Methoden auswählen, Lösungen begründen und das Vorgehen reflektieren.',
]


class WorkingCopyExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = Document(REPO / 'downloads' / f'{STEM}.docx')
        cls.pdf = PdfReader(REPO / 'downloads' / f'{STEM}.pdf')
        cls.web = html.fromstring((REPO / 'chemie-curriculum-arbeitskopie.html').read_text())
        cls.text = ''.join(cls.doc._element.xpath('//w:t/text()'))
        cls.printed = '\n'.join(page.extract_text() or '' for page in cls.pdf.pages)

    def test_one_canonical_afb_definition_in_both_formats(self):
        printed = ''.join(self.printed.split())
        for definition in DEFINITIONS:
            self.assertEqual(self.text.count(definition), 1)
            self.assertEqual(printed.count(''.join(definition.split())), 1)
        self.assertEqual(self.text.count('Anforderungsbereiche: Für die aufgabenbezogene Zuordnung'), 1)

    def test_all_general_and_chemistry_grade_formulations(self):
        tables = self.web.get_element_by_id('abiturpruefung').getparent().xpath('.//table')
        self.assertEqual(len(tables), 2)
        printed = ''.join(self.printed.split())
        for table in tables:
            rows = table.xpath('./tbody/tr')
            self.assertEqual(len(rows), 16)
            for row in rows:
                expected = row[-1].text_content()
                self.assertIn(expected, self.text)
                self.assertIn(''.join(expected.split()), printed)

    def test_native_appendices_and_no_downloads_or_interactive_ui(self):
        headings = [p.text for p in self.doc.paragraphs if p.style.name == 'Heading 1']
        appendices = [heading for heading in headings if heading.startswith('Anhang ')]
        self.assertEqual(len(appendices), 4)
        self.assertEqual(headings[-1], '18 Quellen und Geltungsstand')
        self.assertIn('Aufgabenblatt', self.text)
        self.assertIn('Das folgende Ausgabeblatt', self.text)
        self.assertIn('Klausurergebnis und Rückmeldung', self.text)
        self.assertGreaterEqual(len(self.doc.tables), 66)
        self.assertFalse(self.doc._element.xpath('//w:sdt'))
        self.assertFalse(self.pdf.get_fields())
        for forbidden in ('download', 'herunterladen', 'Vorlage erstellen', 'Vorlage kopieren',
                          'Für die Vorlage angenommene Punktstufe', 'Hier im Fenster öffnen',
                          'Schulische Arbeitshilfe, keine neue Prüfungsordnung'):
            self.assertNotIn(forbidden.lower(), self.text.lower())
        for rel in self.doc.part.rels.values():
            if rel.is_external:
                self.assertNotIn('downloads/', rel.target_ref)
        for page in self.pdf.pages:
            for ref in page.get('/Annots', []):
                uri = ref.get_object().get('/A', {}).get('/URI', '')
                self.assertNotIn('downloads/', uri)

    def test_navigation_references_and_local_step_numbering(self):
        names = self.doc._element.xpath('//w:bookmarkStart/@w:name')
        self.assertEqual(len(names), len(set(names)))
        self.assertTrue(set(self.doc._element.xpath('//w:hyperlink/@w:anchor')).issubset(set(names)))
        self.assertIn('VV 14.2.1 bis 14.3.3', self.text)
        self.assertIn('Kapitel 2.5 und 17', self.text)
        self.assertIn('Didaktisch gefasstes Anwendungsbeispiel in Kapitel 17.', self.text)
        self.assertIn('der individuelle Gesamtaufwand folgt 14.1.', self.text)
        self.assertIn('14.2 bis 14.6', self.text)
        steps = next(p for p in self.doc.paragraphs if p.text.startswith('Vor Beginn: Präsentationsbogen'))
        identity = steps._p.xpath('./w:pPr/w:numPr/w:numId/@w:val')[0]
        start = self.doc.part.numbering_part.element.xpath(
            f'./w:num[@w:numId="{identity}"]/w:lvlOverride/w:startOverride/@w:val')
        self.assertEqual(start, ['1'])
        chapters = [p.text for p in self.doc.paragraphs if p.style.name == 'Heading 1' and re.match(r'^\d+ ', p.text)]
        self.assertEqual(chapters, [heading.text_content() for heading in self.web.xpath('//main//h2')])

    def test_refresh_does_not_duplicate_or_change_text(self):
        with tempfile.TemporaryDirectory(prefix='chemie-export-test-') as folder:
            output = Path(folder) / 'refresh.docx'
            subprocess.run([sys.executable, str(REPO / 'scripts/export-chemie-working-copy.py'),
                            '--output', str(output)], check=True, capture_output=True)
            refreshed = Document(output)
            self.assertEqual(''.join(refreshed._element.xpath('//w:t/text()')), self.text)
            self.assertEqual(len(refreshed.tables), len(self.doc.tables))
            self.assertEqual(len(refreshed._element.xpath('//w:bookmarkStart')),
                             len(self.doc._element.xpath('//w:bookmarkStart')))

    def test_gkl_has_contiguous_examples_and_no_continuation_titles(self):
        gkl = self.web.get_element_by_id('gkl').getparent()
        web_titles = [e.text_content() for e in gkl.xpath('./h3')]
        self.assertEqual([re.match(r'^14\.(\d+)',t).group(1) for t in web_titles],
                         [str(i) for i in range(1,13)])
        self.assertTrue(all('Beispiel ' in t for t in web_titles[1:6]))
        for i in range(2,7):
            heading = gkl.xpath(f'./h3[starts-with(text(),"14.{i} ")]')[0]
            self.assertTrue(heading.getnext().text_content().startswith('Zuordnung:'))
            self.assertIn(heading.getnext().text_content(),self.text)
            self.assertIn(heading.getnext().getnext().text_content(),self.text)
        standalone = html.fromstring((REPO/'chemie-gkl.html').read_text())
        standalone_titles = [e.text_content() for e in standalone.xpath('//main//h2')
                             if re.match(r'^14\.',e.text_content())]
        self.assertEqual(standalone_titles,web_titles)
        word_titles = [p.text for p in self.doc.paragraphs if p.style.name=='Heading 2'
                       and re.match(r'^14\.',p.text)]
        self.assertEqual(word_titles,web_titles)
        for titles in ([e.text_content() for e in self.web.xpath('//main//h2|//main//h3')],
                       [p.text for p in self.doc.paragraphs if p.style.name.startswith('Heading')]):
            self.assertFalse(any(re.search(r'\bFortsetzung\s*$',t) for t in titles))
        self.assertNotIn('Fortsetzung',self.printed)
        self.assertEqual(sum(p.text.startswith('Zuordnung:') for p in self.doc.paragraphs),5)
        self.assertIn('Präsentationsbogen (14.9) oder Experimentierbogen (14.10)',self.text)
        self.assertIn('Rollenmodul aus 14.11',self.text)

    def test_standalone_gkl_exports_match_the_working_copy_structure(self):
        doc = Document(REPO/'downloads/Chemie_GKL_Aufgaben_und_Bewertung.docx')
        pdf = PdfReader(REPO/'downloads/Chemie_GKL_Aufgaben_und_Bewertung.pdf')
        texts = [p.text for p in doc.paragraphs]
        titles = [p.text for p in doc.paragraphs if p.style.name=='Heading 2' and re.match(r'^14\.',p.text)]
        expected = [e.text_content() for e in self.web.get_element_by_id('gkl').getparent().xpath('./h3')]
        self.assertEqual(titles,expected)
        self.assertEqual(sum(t.startswith('Zuordnung:') for t in texts),5)
        self.assertFalse(any('Fortsetzung' in t for t in texts))
        self.assertFalse(any('Fortsetzung' in (p.extract_text() or '') for p in pdf.pages))
        names = doc._element.xpath('//w:bookmarkStart/@w:name')
        self.assertEqual(len(names),len(set(names)))
        self.assertTrue(set(doc._element.xpath('//w:hyperlink/@w:anchor')).issubset(names))
        self.assertEqual(len(doc.tables),20)


if __name__ == '__main__':
    unittest.main()
