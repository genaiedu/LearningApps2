"""Create eight editable A4 Word examples with embedded fonts and quiet styles."""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from lxml import etree
from uuid import uuid5, NAMESPACE_URL
from hashlib import sha256
from struct import unpack_from

DEST=Path(__file__).resolve().parents[1]/'downloads'/'typografie'
DEST.mkdir(parents=True,exist_ok=True)
FONT=DEST.parents[1]/'fonts'/'typografie'/'cardo'

def embed_fonts(d,family_name='Cardo',folder=FONT,files=None,license_file='OFL.txt'):
    """Embed full original fonts; preserve the OFL within every DOCX."""
    files=files or [('Regular','Cardo-Regular.ttf'),('Bold','Cardo-Bold.ttf'),('Italic','Cardo-Italic.ttf')]
    table=d.part.part_related_by(RT.FONT_TABLE);xml=parse_xml(table.blob)
    family=OxmlElement('w:font');family.set(qn('w:name'),family_name)
    for tag,val in [('family','swiss' if family_name=='Source Sans 3' else 'roman'),('pitch','variable')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);family.append(e)
    for variant,file in files:
        raw=(folder/file).read_bytes();entries=unpack_from('>H',raw,4)[0]
        offset=next(unpack_from('>I',raw,20+i*16)[0] for i in range(entries) if raw[12+i*16:16+i*16]==b'OS/2')
        fs_type=unpack_from('>H',raw,offset+8)[0]
        assert fs_type==0 or (fs_type&8 and not fs_type&2), 'Editable font embedding not allowed'
        guid=uuid5(NAMESPACE_URL,sha256(raw).hexdigest());key=bytes.fromhex(guid.hex)[::-1]
        data=bytearray(raw)
        for i in range(32):data[i]^=key[i%16]
        part=Part(PackURI('/word/fonts/'+Path(file).stem+'.odttf'),
                  'application/vnd.openxmlformats-officedocument.obfuscatedFont',bytes(data),d.part.package)
        e=OxmlElement('w:embed'+variant);e.set(qn('r:id'),table.relate_to(part,RT.FONT))
        e.set(qn('w:fontKey'),'{'+str(guid).upper()+'}');e.set(qn('w:subsetted'),'0');family.append(e)
    xml.append(family);table._blob=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    licence=etree.Element('{urn:typografie-atelier:license}font-license',font=family_name)
    licence.text=(folder/license_file).read_text()
    part=Part(PackURI('/customXml/'+family_name.lower().replace(' ','')+'-license.xml'),'application/xml',etree.tostring(licence,encoding='UTF-8'),d.part.package)
    d.part.relate_to(part,RT.CUSTOM_XML)
    if d.settings.element.find(qn('w:embedTrueTypeFonts')) is None:
        d.settings.element.append(OxmlElement('w:embedTrueTypeFonts'))

def base(title):
    d=Document(); sec=d.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.left_margin=Mm(32);sec.right_margin=Mm(32);sec.top_margin=Mm(27);sec.bottom_margin=Mm(35)
    sec.header_distance=Mm(12);sec.footer_distance=Mm(15)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=d.styles[name];st.font.color.rgb=RGBColor(0,0,0)
        st.font.name='Cardo'
        st.font.size=Pt({'Normal':12,'Title':22,'Subtitle':12,'Heading 1':12,'Heading 2':12}[name])
        st.font.bold=name=='Heading 1';st.font.italic=name in ['Subtitle','Heading 2']
        st.paragraph_format.space_after=Pt(0)
        st.paragraph_format.line_spacing=1.22
        st.paragraph_format.widow_control=True
        st.paragraph_format.first_line_indent=Mm(0)
    d.styles['Title'].font.bold=False
    # The bundled starter can contain decorative title rules and theme fonts.
    # Remove them rather than inheriting an unrelated document design.
    for st in d.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        for rfonts in st.element.iter(qn('w:rFonts')):
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                rfonts.attrib.pop(qn('w:'+attr),None)
            for attr in ['ascii','hAnsi','eastAsia','cs']:rfonts.set(qn('w:'+attr),'Cardo')
    normal=d.styles['Normal'];normal.paragraph_format.first_line_indent=Mm(4.23)
    lang=normal.element.get_or_add_rPr().find(qn('w:lang'))
    if lang is None:lang=OxmlElement('w:lang');normal.element.get_or_add_rPr().append(lang)
    lang.set(qn('w:val'),'de-DE')
    for lang in d.styles.element.iter(qn('w:lang')):lang.set(qn('w:val'),'de-DE')
    if title.startswith('Kurzbericht'):normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    first=d.styles.add_style('Erster Absatz',WD_STYLE_TYPE.PARAGRAPH)
    first.base_style=normal;first.paragraph_format.first_line_indent=Mm(0);first.next_paragraph_style=normal
    for name in ['Heading 1','Heading 2']:
        d.styles[name].next_paragraph_style=first
        d.styles[name].paragraph_format.space_after=Pt(8)
    d.styles['Title'].paragraph_format.space_after=Pt(14)
    d.styles['Title'].paragraph_format.keep_with_next=True
    d.styles['Subtitle'].paragraph_format.space_after=Pt(16)
    d.styles['Subtitle'].paragraph_format.keep_with_next=True
    foot=d.styles['Footnote Text'] if 'Footnote Text' in d.styles else d.styles.add_style('Footnote Text',WD_STYLE_TYPE.PARAGRAPH)
    foot.base_style=normal;foot.font.name='Cardo';foot.font.size=Pt(10)
    foot.paragraph_format.line_spacing=1.22;foot.paragraph_format.space_after=Pt(0)
    for tag,val in [('autoHyphenation','1'),('consecutiveHyphenLimit','2'),('doNotHyphenateCaps','1')]:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),val);d.settings.element.append(el)
    d.styles['Heading 1'].paragraph_format.space_before=Pt(18)
    d.styles['Heading 2'].paragraph_format.space_before=Pt(12)
    d.styles['Heading 1'].paragraph_format.keep_with_next=True
    d.styles['Heading 2'].paragraph_format.keep_with_next=True
    d.core_properties.title=title;d.core_properties.author='Typografie-Atelier · Dr. Claus Unterberg'
    d.core_properties.subject='Klassisch orientierte DIN-A4-Vorlage; Cardo vollständig eingebettet; Platzhalter ersetzen.'
    d.core_properties.comments='Eigene zeitgenössische Gestaltung nach Prinzipien des klassischen Buchsatzes, keine Originalvorlage Tschicholds. Keine DIN-5008-Zertifizierung. Cardo unter SIL OFL 1.1; Lizenztext im Dokumentpaket. Druckvorschau vor Weitergabe prüfen.'
    d.core_properties.keywords='Typografie, Vorlage, DIN A4'
    embed_fonts(d)
    return d

def p(d,text,style=None):
    if style is None and d.paragraphs and d.paragraphs[-1].style.name in ['Title','Subtitle','Heading 1','Heading 2']:
        style='Erster Absatz'
    return d.add_paragraph(text,style)
def footer(d):
    r=d.sections[0].footer.paragraphs[0];r.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r.paragraph_format.first_line_indent=Mm(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    run=OxmlElement('w:r');rp=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'),'Cardo');rf.set(qn('w:hAnsi'),'Cardo');rp.append(rf)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24');rp.append(sz);run.append(rp)
    text=OxmlElement('w:t');text.text='1';run.append(text);field.append(run);r._p.append(field)
def modernize(d):
    family='Source Sans 3'
    for st in d.styles:
        for rf in st.element.iter(qn('w:rFonts')):
            for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),family)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name=family;st.font.size=Pt(16 if name=='Title' else 12)
        st.font.bold=name in ['Title','Heading 1'];st.font.italic=name=='Heading 2'
    normal=d.styles['Normal'];normal.paragraph_format.first_line_indent=Mm(0)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.3
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    # Only the separate title block has a larger grade; whitespace separates it.
    d.styles['Title'].paragraph_format.space_after=Pt(18)
    d.styles['Subtitle'].paragraph_format.space_after=Pt(22)
    for sec in d.sections:
        for rf in sec.footer._element.iter(qn('w:rFonts')):
            for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),family)
    table=d.part.part_related_by(RT.FONT_TABLE);xml=parse_xml(table.blob)
    for f in list(xml):
        if f.get(qn('w:name'))=='Cardo':
            for e in f:
                rid=e.get(qn('r:id'))
                if rid:table.drop_rel(rid)
            xml.remove(f)
    table._blob=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    for rid,rel in list(d.part.rels.items()):
        if not rel.is_external and str(rel.target_part.partname)=='/customXml/cardo-license.xml':d.part.drop_rel(rid)
    static=FONT.parent/'sourcesans3'/'static'
    embed_fonts(d,family,static,[('Regular','SourceSans3-Regular.otf'),('Bold','SourceSans3-Bold.otf'),('Italic','SourceSans3-It.otf'),('BoldItalic','SourceSans3-BoldIt.otf')],'LICENSE.txt')
    d.core_properties.title+=' · moderne Variante'
    d.core_properties.subject='Moderne A4-Vorlage; Source Sans 3 vollständig eingebettet; ein Schriftgrad im Lesetext.'
    d.core_properties.comments='Eigene moderne Gebrauchstypografie: humanistische Sans, linksbündiger Satz, Absatzabstände. Größerer Titel nur im abgesetzten Kopfbereich. Source Sans 3 unter SIL OFL 1.1; vollständiger Lizenztext im Dokumentpaket.'

def save(d,name):
    d.save(DEST/(name+'.docx'));print(name+'.docx')
    modern=Document(DEST/(name+'.docx'));modernize(modern)
    modern.save(DEST/(name+'-modern.docx'));print(name+'-modern.docx')

d=base('Sachlicher Brief · A4')
d.styles['Normal'].paragraph_format.first_line_indent=Mm(0)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Heading 1'].font.size=Pt(12)
d.styles['Heading 1'].font.bold=True
d.styles['Heading 1'].paragraph_format.space_after=Pt(12)
p(d,'[Vorname Nachname] · [Straße und Hausnummer] · [Postleitzahl Ort]','Subtitle')
p(d,'[Name der Empfängerin / des Empfängers]\n[Organisation]\n[Straße und Hausnummer]\n[Postleitzahl Ort]')
q=p(d,'[Ort], [Datum]');q.alignment=WD_ALIGN_PARAGRAPH.RIGHT
p(d,'[Aussagekräftiger Betreff]','Heading 1')
p(d,'Sehr geehrte [Anrede und Name],')
p(d,'[Nennen Sie im ersten Absatz den Anlass und Ihr Anliegen. Die lesende Person soll sofort erkennen, worum es geht.]')
p(d,'[Begründen Sie Ihr Anliegen im zweiten Absatz. Ordnen Sie Informationen so, dass ein Gedanke auf den nächsten aufbaut. Verwenden Sie kurze, sachliche Sätze.]')
p(d,'[Formulieren Sie zum Schluss den gewünschten nächsten Schritt und gegebenenfalls einen Termin. Lassen Sie Raum für eine freundliche Antwort.]')
p(d,'Mit freundlichen Grüßen')
q=p(d,'[Unterschrift]\n[Vorname Nachname]');q.paragraph_format.space_before=Pt(20)
q=p(d,'[Anlagen, falls vorhanden]','Subtitle');q.paragraph_format.space_before=Pt(10)
save(d,'brief-a4')

d=base('Kurzbericht · A4')
p(d,'[Titel des Kurzberichts]','Title')
p(d,'[Name] · [Kurs / Organisation] · [Datum]','Subtitle')
p(d,'Fragestellung','Heading 1')
p(d,'[Welche Frage soll der Bericht beantworten? Formulieren Sie Gegenstand, Anlass und Ziel in zwei bis vier Sätzen. Die Überschriften dieser Vorlage sind echte Formatvorlagen, keine manuell fett gesetzten Absätze.]')
p(d,'Vorgehen','Heading 1')
p(d,'[Beschreiben Sie, welche Materialien, Beobachtungen oder Quellen Sie verwendet haben. Trennen Sie eigene Feststellungen von Informationen anderer Personen. Fügen Sie Quellenbelege dort ein, wo sie die Aussage stützen.]')
p(d,'Ergebnisse','Heading 1')
p(d,'[Stellen Sie die wichtigsten Ergebnisse in nachvollziehbarer Reihenfolge dar. Beginnen Sie pro Absatz einen neuen Gedanken. Ein kurzer Bericht braucht nicht mehrere Schriftfamilien: Eine Leseschrift und eine klar abgestufte Überschriftenstruktur reichen häufig aus.]')
p(d,'Einordnung und Ausblick','Heading 1')
p(d,'[Beantworten Sie die Ausgangsfrage. Nennen Sie Grenzen Ihrer Untersuchung und einen sinnvollen nächsten Schritt. Prüfen Sie vor der Weitergabe, ob alle Platzhalter ersetzt sind.]')
p(d,'Quellen','Heading 1')
p(d,'[Autor/in oder Institution: Titel, Jahr; bei Internetquellen URL und Abrufdatum. Verwenden Sie für alle Einträge dasselbe Schema.]')
footer(d);save(d,'kurzbericht-a4')

d=base('Einladung · A4')
d.sections[0].top_margin=Mm(38)
q=p(d,'Einladung','Subtitle');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p(d,'[Ein besonderer Anlass]','Title');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p(d,'[Tag und Datum] · [Uhrzeit]','Subtitle');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p(d,'[Veranstaltungsort und Adresse]','Erster Absatz');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(26)
p(d,'Liebe [Name / Gäste],','Erster Absatz')
q=p(d,'[Sprechen Sie Ihre Gäste persönlich an. Beschreiben Sie in einem kurzen Absatz, worauf sie sich freuen dürfen und warum Sie diesen Anlass gemeinsam feiern möchten.]','Erster Absatz');q.paragraph_format.space_before=Pt(10)
p(d,'[Ergänzen Sie die wichtigsten praktischen Hinweise: Ablauf, Anreise, Zugänglichkeit oder Dinge, die mitgebracht werden sollen. Zentrieren Sie nicht jeden längeren Absatz; linksbündiger Text lässt sich leichter verfolgen.]')
p(d,'[Bitte um Rückmeldung bis Datum]','Heading 2')
p(d,'[Kontakt für die Rückmeldung]')
q=p(d,'Wir freuen uns auf Sie / euch!','Erster Absatz');q.paragraph_format.space_before=Pt(22)
q=p(d,'[Name der einladenden Person]','Erster Absatz');q.paragraph_format.space_before=Pt(10)
save(d,'einladung-a4')

d=base('Ergebnisprotokoll · A4')
p(d,'[Sitzung / Projektgruppe]','Title')
p(d,'Ergebnisprotokoll','Subtitle')
p(d,'Datum und Zeit: [Angaben]\nOrt: [Angabe]\nTeilnehmende: [Namen]\nProtokoll: [Name]')
p(d,'1. Anlass und Tagesordnung','Heading 1')
p(d,'[Nennen Sie den Anlass und die behandelten Themen. Dieses Ergebnisprotokoll hält Beschlüsse fest; es ist keine wörtliche Mitschrift.]')
p(d,'2. Ergebnisse und Beschlüsse','Heading 1')
p(d,'[Thema 1]','Heading 2')
p(d,'[Ergebnis oder Beschluss, gegebenenfalls mit Begründung und abweichenden Positionen.]')
p(d,'[Thema 2]','Heading 2')
p(d,'[Ergebnis oder Beschluss. Machen Sie kenntlich, wenn eine Frage noch offen ist.]')
p(d,'3. Nächste Schritte','Heading 1')
p(d,'Aufgabe: [konkrete Tätigkeit]\nVerantwortlich: [Name]\nTermin: [Datum]')
p(d,'4. Nächster Termin','Heading 1')
p(d,'[Datum, Uhrzeit und Ort; erforderliche Vorbereitung]')
footer(d);save(d,'ergebnisprotokoll-a4')
