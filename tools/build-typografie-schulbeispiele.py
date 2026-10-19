"""Twelve editable examples. Legacy and Abitur templates stay untouched; shared DIV-12 geometry."""
from pathlib import Path
from copy import deepcopy
import hashlib,json,subprocess,uuid,struct
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement,parse_xml
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
DEST=ROOT/'downloads/typografie'
PAIRS=[('Regular',400,False),('Italic',400,True),('Bold',700,False),('BoldItalic',700,True)]
FAMILIES={'EB Garamond':ROOT/'fonts/typografie/ebgaramond/static','Fira Sans':ROOT/'fonts/typografie/firasans'}
PREFIX={'EB Garamond':'EBGaramond','Fira Sans':'FiraSans'}
GEOMETRY=json.loads(subprocess.check_output(['node','-e',"const S=require('./scripts/typografie-satzspiegel-core.js');console.log(JSON.stringify({single:S.geometry('a4','raster12','single'),spread:S.geometry('a4','raster12')}))"],cwd=ROOT,text=True))

def prop(parent,tag,val=None,attrs=None):
    node=OxmlElement(tag)
    if val is not None:node.set(qn('w:val'),str(val))
    for key,value in (attrs or {}).items():node.set(qn(key),str(value))
    parent.append(node);return node

def fonts(d):
    table=d.part.part_related_by(RT.FONT_TABLE)
    xml=parse_xml(table.blob)
    for child in list(xml):xml.remove(child)
    for family,folder in FAMILIES.items():
        f=prop(xml,'w:font',attrs={'w:name':family});prop(f,'w:family','roman' if family=='EB Garamond' else 'swiss')
        for cut,weight,italic in PAIRS:
            path=folder/(PREFIX[family]+'-'+cut+'.ttf');raw=path.read_bytes()
            count=struct.unpack_from('>H',raw,4)[0]
            tables={raw[12+i*16:16+i*16]:struct.unpack_from('>II',raw,20+i*16) for i in range(count)}
            assert b'fvar' not in tables, 'Only official static TTFs may be embedded'
            rights=struct.unpack_from('>H',raw,tables[b'OS/2'][0]+8)[0];assert rights==0
            guid=uuid.uuid5(uuid.NAMESPACE_URL,hashlib.sha256(raw).hexdigest());key=guid.bytes[::-1];encoded=bytearray(raw)
            for i in range(32):encoded[i]^=key[i%16]
            part=Part(PackURI('/word/fonts/'+path.stem+'.odttf'),'application/vnd.openxmlformats-officedocument.obfuscatedFont',bytes(encoded),d.part.package)
            prop(f,'w:embed'+cut,attrs={'r:id':table.relate_to(part,RT.FONT),'w:fontKey':'{'+str(guid).upper()+'}','w:subsetted':'0'})
        license_path=folder.parent/'OFL.txt' if family=='EB Garamond' else folder/'OFL.txt'
        license_xml=etree.Element('{urn:typografie-atelier:license}font-license',font=family);license_xml.text=license_path.read_text()
        part=Part(PackURI('/customXml/'+PREFIX[family]+'-OFL.xml'),'application/xml',etree.tostring(license_xml,encoding='UTF-8'),d.part.package);d.part.relate_to(part,RT.CUSTOM_XML)
    table._blob=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    prop(d.settings.element,'w:embedTrueTypeFonts');prop(d.settings.element,'w:saveSubsetFonts','0')

def style(d,name,family,size=12,bold=False,italic=False,before=0,after=0,indent=0,align=None,keep=False):
    st=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    if name!='Normal':st.base_style=d.styles['Normal']
    st.font.name=family;st.font.size=Pt(size);st.font.bold=bold;st.font.italic=italic;st.font.color.rgb=RGBColor(0,0,0)
    rpr=st.element.get_or_add_rPr();rf=rpr.get_or_add_rFonts()
    for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),family)
    for attr in list(rf.attrib):
        if 'Theme' in attr:del rf.attrib[attr]
    prop(rpr,'w:lang',attrs={'w:val':'de-DE'})
    prop(rpr,'w14:ligatures',attrs={'w14:val':'standard'})
    data=name in ['Header','Footer','Caption','Tabelleninhalt','Punkte','Rücksendezeile','Anschrift','Metadaten']
    prop(rpr,'w14:numForm',attrs={'w14:val':'lining' if data else 'oldStyle'})
    prop(rpr,'w14:numSpacing',attrs={'w14:val':'tabular' if data else 'proportional'})
    pf=st.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=1.3 if family=='Fira Sans' else 1.25
    pf.first_line_indent=Pt(indent);pf.widow_control=True;pf.keep_with_next=keep
    if align is not None:pf.alignment=align
    return st

def base(title,modern=False,spread=False,brief=False):
    d=Document();s=d.sections[0];g=GEOMETRY['spread' if spread else 'single'];body='Fira Sans' if modern else 'EB Garamond'
    s.page_width=Mm(g['width']);s.page_height=Mm(g['height']);s.left_margin=Mm(g['inner']);s.right_margin=Mm(g['outer']);s.top_margin=Mm(g['top']);s.bottom_margin=Mm(g['bottom']);s.gutter=Mm(0);s.header_distance=Mm(11);s.footer_distance=Mm(18)
    if spread:prop(d.settings.element,'w:mirrorMargins')
    # Strip theme and decorative rules from the starter, including unused default styles.
    for st in d.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        for rf in st.element.iter(qn('w:rFonts')):
            for a in list(rf.attrib):del rf.attrib[a]
            for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),body)
    normal=style(d,'Normal',body,indent=0 if modern or brief else 12,after=7 if modern or brief else 0,align=WD_ALIGN_PARAGRAPH.LEFT if modern or brief else WD_ALIGN_PARAGRAPH.JUSTIFY)
    style(d,'Erster Absatz',body,after=7 if modern or brief else 0,align=normal.paragraph_format.alignment)
    for name,bold,italic in [('Heading 1',True,False),('Heading 2',False,True)]:
        st=style(d,name,body,bold=bold,italic=italic,before=15,after=7,keep=True);st.next_paragraph_style=d.styles['Erster Absatz']
    style(d,'Title',body,22 if not modern else 18,bold=True,after=15,keep=True)
    style(d,'Subtitle',body,12,italic=True,after=14,keep=True)
    for name,size in [('Header',9),('Footer',9),('Caption',10),('Tabelleninhalt',11),('Metadaten',10),('Punkte',10)]:style(d,name,'Fira Sans',size,after=5,align=WD_ALIGN_PARAGRAPH.LEFT)
    style(d,'Aufgabe',body,bold=True,before=10,after=5,keep=True)
    style(d,'Aufgabentext',body,after=7,align=WD_ALIGN_PARAGRAPH.LEFT)
    style(d,'Material',body,after=8,align=WD_ALIGN_PARAGRAPH.LEFT if modern else WD_ALIGN_PARAGRAPH.JUSTIFY)
    style(d,'Antwortraum',body,size=12,after=6)
    d.styles['Antwortraum'].paragraph_format.line_spacing=Pt(20)
    for name in ['Aufgabe','Aufgabentext']:
        d.styles[name].paragraph_format.left_indent=Mm(7);d.styles[name].paragraph_format.first_line_indent=Mm(-7 if name=='Aufgabe' else 0)
    points=d.styles['Punkte'].paragraph_format;points.right_indent=Mm(-17);points.tab_stops.add_tab_stop(Mm(g['textWidth']+14),WD_TAB_ALIGNMENT.RIGHT)
    style(d,'Briefdatum','Fira Sans',10,after=10,align=WD_ALIGN_PARAGRAPH.RIGHT)
    style(d,'Einladungskopf',body,12,italic=True,after=14,align=WD_ALIGN_PARAGRAPH.CENTER,keep=True)
    style(d,'Einladungstitel',body,22 if not modern else 18,bold=True,after=16,align=WD_ALIGN_PARAGRAPH.CENTER,keep=True)
    style(d,'Einladungsangaben','Fira Sans',11,after=14,align=WD_ALIGN_PARAGRAPH.CENTER)
    style(d,'Rücklauf',body,11,before=12,after=6)
    borders=prop(d.styles['Rücklauf'].element.get_or_add_pPr(),'w:pBdr');prop(borders,'w:top',attrs={'w:val':'dashed','w:sz':'4','w:space':'8','w:color':'000000'})
    for name,y,h,size in [('Rücksendezeile',45,17.7,8),('Anschrift',62.7,27.3,11)]:
        st=style(d,name,'Fira Sans',size,align=WD_ALIGN_PARAGRAPH.LEFT)
        prop(st.element.get_or_add_pPr(),'w:framePr',attrs={'w:w':str(round(80*1440/25.4)),'w:h':str(round(h*1440/25.4)),'w:x':str(round(25*1440/25.4)),'w:y':str(round(y*1440/25.4)),'w:hAnchor':'page','w:vAnchor':'page','w:wrap':'notBeside','w:hRule':'atLeast'})
    style(d,'Briefbetreff',body,bold=True,before=207,after=12,keep=True)
    for tag,val in [('autoHyphenation','1'),('consecutiveHyphenLimit','2'),('doNotHyphenateCaps','1')]:prop(d.settings.element,'w:'+tag,val)
    d.core_properties.title=title+(' · modern' if modern else ' · klassisch');d.core_properties.author='Typografie-Atelier · Dr. Claus Unterberg'
    d.core_properties.subject='Eigenes Beispiel · Zwölferteilung · EB Garamond und Fira Sans vollständig statisch eingebettet'
    d.core_properties.comments='Keine DIN-5008-Zertifizierung. Beide SIL-OFL-1.1-Lizenztexte befinden sich im Dokumentpaket. Seitenansicht nach Änderungen prüfen.'
    fonts(d);return d

def add(d,text,name='Normal'):
    p=d.add_paragraph(text,name)
    # Word paragraph frames need the positional record on the anchored paragraph.
    # Keep its geometry centralized in the named style; clone only the layout record.
    frame=d.styles[name].element.find('w:pPr/w:framePr',{'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
    if frame is not None:p._p.get_or_add_pPr().append(deepcopy(frame))
    return p
def field(p,instruction):
    f=prop(p._p,'w:fldSimple',attrs={'w:instr':instruction});r=prop(f,'w:r');t=prop(r,'w:t');t.text='1'
def footer(d):
    p=d.sections[0].footer.paragraphs[0];p.style='Footer';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Seite ');field(p,'PAGE');p.add_run(' von ');field(p,'NUMPAGES')
def save(d,key,modern):
    target=DEST/(key+'-a4-zw12'+('-modern' if modern else '')+'.docx');d.save(target);print(target.name)

for key in ['brief','kurzbericht','einladung','ergebnisprotokoll']:
    source=Document(DEST/(key+'-a4.docx'))
    for modern in [False,True]:
        d=base(source.core_properties.title,modern,spread=key in ['kurzbericht','ergebnisprotokoll'],brief=key=='brief')
        for index,p in enumerate(source.paragraphs):
            if not p.text:continue
            name=p.style.name
            if key=='brief':
                if index==0:name='Rücksendezeile'
                elif index==1:name='Anschrift'
                elif index==2:name='Briefdatum'
                elif index==3:name='Briefbetreff'
            elif key=='einladung' and index<4:name=['Einladungskopf','Einladungstitel','Einladungsangaben','Einladungsangaben'][index]
            elif key=='ergebnisprotokoll' and '\n' in p.text:name='Metadaten'
            add(d,p.text,name if name in d.styles else 'Normal')
        footer(d);save(d,key,modern)

for modern in [False,True]:
    d=base('Arbeitsblatt · Geschichte',modern)
    add(d,'Fach: [Fach]     Name: [Name]     Datum: [Datum]','Metadaten')
    d.sections[0].header.paragraphs[0].text='[Fach] · Arbeitsblatt';d.sections[0].header.paragraphs[0].style='Header'
    add(d,'[Thema des Arbeitsblatts]','Title')
    add(d,'Material','Heading 1');add(d,'[Hier den Ausgangstext oder die Quelle einsetzen. Material und Arbeitsauftrag bleiben sichtbar getrennt. Quellenangabe ergänzen.]','Material')
    add(d,'[Abbildung bei Bedarf hier einfügen]','Erster Absatz');add(d,'Abbildung 1: [Was zeigt die Abbildung? Quelle ergänzen.]','Caption')
    add(d,'1. Quelle untersuchen','Aufgabe');add(d,'Lies das Material. Nenne die Position der Verfasserin oder des Verfassers und belege sie mit zwei Textstellen.','Aufgabentext')
    add(d,'2. Urteil begründen','Aufgabe');add(d,'Vergleiche die Position mit einer anderen Perspektive. Begründe dein Urteil.','Aufgabentext')
    add(d,'Arbeitszeit: [Dauer] · Hilfsmittel: [Angaben]','Metadaten');footer(d);save(d,'arbeitsblatt',modern)
    d=base('Elternbrief · Beispiel',modern,brief=True)
    p=d.sections[0].header.paragraphs[0];p.style='Header';p.text='[Schulbriefkopf] · [Logo der Schule]'
    add(d,'[Schule] · [Straße] · [Postleitzahl Ort]','Rücksendezeile');add(d,'[Name der Empfängerin / des Empfängers]\n[Straße und Hausnummer]\n[Postleitzahl Ort]','Anschrift')
    add(d,'[Ort], [Datum]','Briefdatum');add(d,'[Anlass des Elternbriefs]','Briefbetreff');add(d,'Liebe Eltern,','Erster Absatz');add(d,'[Nennen Sie den Anlass und die wichtigsten Informationen. Termin, Ort, Kosten und der nächste Schritt müssen klar erkennbar sein.]');add(d,'[Ergänzen Sie die Kontaktmöglichkeit für Rückfragen.]');add(d,'Mit freundlichen Grüßen\n[Name und Funktion]','Erster Absatz')
    add(d,'Rücklaufzettel · bei Bedarf entfernen','Rücklauf');add(d,'Name des Kindes: ____________________\n[ ] Ich habe die Informationen erhalten.\n[ ] Ich bitte um Rücksprache.\nDatum / Unterschrift: ____________________','Aufgabentext');footer(d);save(d,'elternbrief',modern)
